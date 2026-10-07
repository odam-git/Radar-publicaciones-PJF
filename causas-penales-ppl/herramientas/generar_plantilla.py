#!/usr/bin/env python3
"""
Genera plantillas/plantilla_captura_ppl.xlsx: la hoja que llenan los juzgados.

- Hoja "Captura": encabezados exactos que lee la aplicación, con listas desplegables
  (entidad, tipo de procedimiento, etapa procesal, motivo de la privación de libertad, situación de la
  otra causa, revisión de la medida y motivo de baja),
  validación de fechas y de pena, y renglones de ejemplo ficticios en gris.
- Hoja "Instrucciones": qué va en cada columna, si es obligatoria y en qué formato.
- Hoja "Catálogos": las listas que alimentan los desplegables.

Uso:  python3 herramientas/generar_plantilla.py      (requiere: pip install openpyxl)
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "plantillas" / "plantilla_captura_ppl.xlsx"
FILAS = 2000  # renglones con validación

ENTIDADES = [
    "Aguascalientes", "Baja California", "Baja California Sur", "Campeche", "Chiapas", "Chihuahua",
    "Ciudad de México", "Coahuila", "Colima", "Durango", "Estado de México", "Guanajuato", "Guerrero",
    "Hidalgo", "Jalisco", "Michoacán", "Morelos", "Nayarit", "Nuevo León", "Oaxaca", "Puebla",
    "Querétaro", "Quintana Roo", "San Luis Potosí", "Sinaloa", "Sonora", "Tabasco", "Tamaulipas",
    "Tlaxcala", "Veracruz", "Yucatán", "Zacatecas",
]
# Los catálogos se toman del generador de datos para que plantilla, datos y aplicación coincidan.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from generar_datos_ficticios import (  # noqa: E402
    COLUMNAS as COLUMNAS_DATOS, ETAPAS_V4, MOTIVOS, MOTIVOS_BAJA, RESULTADOS_REVISION, SITUACIONES_OTRA as SITUACIONES,
)

# Opciones de captura: las 8 etapas en seguimiento más las dos que sacan el registro del universo.
ETAPAS = ETAPAS_V4 + ["Amparo directo", "Sentencia ejecutoriada"]
TIPOS = ["Ordinario", "Sumario"]
SI_NO = ["Sí", "No"]

# (columna, obligatoria, tipo, descripción)
COLUMNAS = [
    ("ID", "Sí", "texto", "Número consecutivo del renglón."),
    ("ID de persona", "Recomendado", "texto", "Clave única de la persona (la misma en todas sus causas). Permite contar personas y causas por separado."),
    ("Nombre completo", "Sí", "texto", "Nombre(s) y apellidos. Sin alias."),
    ("Expediente", "Sí", "texto", "Número de expediente."),
    ("Causa penal", "Sí", "texto", "Número de causa penal vigente."),
    ("Delito(s)", "Sí", "texto", "Separar varios delitos con punto y coma."),
    ("Entidad Federativa", "Sí", "lista", "Entidad del Juzgado de Distrito. Con la fecha de inicio de la averiguación previa permite verificar el sistema aplicable."),
    ("Circuito", "Sí", "texto", "Por ejemplo: Sexto Circuito."),
    ("Juzgado de Distrito", "Sí", "texto", "Denominación completa del juzgado."),
    ("Lugar de reclusión", "Sí", "texto", "Centro penitenciario."),
    ("Motivo de la privación de libertad", "Sí", "lista", "Elegir de la lista."),
    ("Otra causa: autoridad", "Si aplica", "texto", "Solo si el motivo menciona otra causa: autoridad que la conoce."),
    ("Otra causa: situación", "Si aplica", "lista", "Solo si el motivo menciona otra causa."),
    ("Fecha de inicio de la averiguación previa", "Recomendado", "fecha",
     "DD/MM/AAAA. Si es igual o posterior a la entrada en vigor del Código Nacional en la entidad, la aplicación pide verificar el sistema (transitorio Cuarto, DOF 18-06-2008)."),
    ("Fecha de auto de formal prisión", "Sí", "fecha", "Fecha del PRIMER auto de formal prisión (DD/MM/AAAA). Base del semáforo (art. 161 CFPP)."),
    ("Tipo de procedimiento", "Sí", "lista", "Ordinario o Sumario (art. 152 CFPP). Define el plazo de la instrucción."),
    ("Etapa procesal", "Sí", "lista",
     "Elegir de la lista. 'Sentencia ejecutoriada' da de baja el registro; 'Amparo directo' lo saca del seguimiento (a disposición del Tribunal Colegiado, art. 191 Ley de Amparo)."),
    ("Instancia actual", "Sí", "texto", "Órgano que conoce actualmente (juzgado, tribunal unitario con toca, tribunal colegiado)."),
    ("Fecha de cierre de instrucción", "Si aplica", "fecha", "DD/MM/AAAA. Desde conclusiones en adelante (art. 150 CFPP)."),
    ("Fecha de la audiencia de vista", "Si aplica", "fecha", "DD/MM/AAAA. Desde la audiencia de vista en adelante (arts. 305 y 97 CFPP)."),
    ("Último acto procesal", "Sí", "texto", "Reseña breve de la última actuación."),
    ("Fecha del último acto procesal", "Sí", "fecha", "DD/MM/AAAA."),
    ("Fecha de sentencia de primera instancia", "Si aplica", "fecha", "Solo si ya hay sentencia de primera instancia (no firme)."),
    ("Pena impuesta (no firme)", "Si aplica", "número", "En años; se admiten decimales (12.5). Solo se muestra como dato informativo."),
    ("Revisión de la medida: solicitada", "Sí", "lista", "Sí o No (quinto transitorio, DOF 17-06-2016)."),
    ("Revisión de la medida: fecha", "Si aplica", "fecha", "DD/MM/AAAA de la solicitud o resolución."),
    ("Revisión de la medida: resultado", "Si aplica", "lista", "Elegir de la lista."),
    ("Fecha de ejecutoria", "Si aplica", "fecha", "Llenar solo si la sentencia causó ejecutoria (art. 360 CFPP); el registro saldrá del seguimiento."),
    ("Motivo de baja", "Si aplica", "lista", "Solo si el registro sale del seguimiento; elegir de la lista."),
    ("Fecha de corte", "Sí", "fecha", "Fecha en que se llenó la información (la misma para todos los renglones)."),
    ("Observaciones", "No", "texto", "Información adicional relevante."),
]
assert [c[0] for c in COLUMNAS] == COLUMNAS_DATOS, "Las columnas de la plantilla deben coincidir con las del generador"

EJEMPLO = [
    "1", "PER-0001", "Persona Ficticia Ejemplo", "EXP-FIC-0001/2010", "CP-FIC-001/2010",
    "Delito genérico de ejemplo", "Puebla", "Sexto Circuito", "Juzgado Primero de Distrito en Materia Penal en Puebla",
    "Centro de Reinserción Ficticio", "Solo por esta causa federal", "", "", "02/02/2010", "15/03/2010", "Ordinario",
    "Instrucción", "Juzgado Primero de Distrito en Materia Penal en Puebla", "", "", "Se desahogó una prueba testimonial.",
    "10/09/2026", "", "", "No", "", "", "", "", "01/10/2026", "Renglón de ejemplo: bórrelo antes de capturar.",
]
assert len(EJEMPLO) == len(COLUMNAS)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Captura"
    cab = PatternFill("solid", fgColor="0F1E33")
    for i, (nombre, oblig, _, desc) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=1, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF", size=12)
        c.fill = cab
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(16, min(60, len(nombre) + 8))
    for i, v in enumerate(EJEMPLO, start=1):
        c = ws.cell(row=2, column=i, value=v)
        c.font = Font(italic=True, color="5B6573")
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "A2"

    # Catálogos
    cat = wb.create_sheet("Catálogos")
    catalogos = [("Entidades", ENTIDADES), ("Etapas", ETAPAS), ("Motivos", MOTIVOS), ("Situaciones", SITUACIONES),
                 ("Tipo de procedimiento", TIPOS), ("Sí / No", SI_NO), ("Resultado de la revisión", RESULTADOS_REVISION),
                 ("Motivo de baja", MOTIVOS_BAJA)]
    for col, (titulo, valores) in enumerate(catalogos, start=1):
        cat.cell(row=1, column=col, value=titulo).font = Font(bold=True)
        for f, v in enumerate(valores, start=2):
            cat.cell(row=f, column=col, value=v)
        cat.column_dimensions[get_column_letter(col)].width = 48

    def lista(columna, letra_cat, n):
        dv = DataValidation(type="list", formula1=f"='Catálogos'!${letra_cat}$2:${letra_cat}${n + 1}",
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="Valor no válido", error="Elija un valor de la lista.")
        ws.add_data_validation(dv)
        dv.add(f"{columna}2:{columna}{FILAS}")

    letra = {nombre: get_column_letter(i) for i, (nombre, *_r) in enumerate(COLUMNAS, start=1)}
    lista(letra["Entidad Federativa"], "A", len(ENTIDADES))
    lista(letra["Etapa procesal"], "B", len(ETAPAS))
    lista(letra["Motivo de la privación de libertad"], "C", len(MOTIVOS))
    lista(letra["Otra causa: situación"], "D", len(SITUACIONES))
    lista(letra["Tipo de procedimiento"], "E", len(TIPOS))
    lista(letra["Revisión de la medida: solicitada"], "F", len(SI_NO))
    lista(letra["Revisión de la medida: resultado"], "G", len(RESULTADOS_REVISION))
    lista(letra["Motivo de baja"], "H", len(MOTIVOS_BAJA))

    for nombre, _, tipo, _ in COLUMNAS:
        rango = f"{letra[nombre]}2:{letra[nombre]}{FILAS}"
        if tipo == "fecha":
            dv = DataValidation(type="date", operator="between", formula1="DATE(1970,1,1)", formula2="DATE(2100,12,31)",
                                allow_blank=True, showErrorMessage=True, showInputMessage=True,
                                promptTitle="Fecha", prompt="Escriba la fecha como DD/MM/AAAA.",
                                errorTitle="Fecha no válida", error="Escriba una fecha válida (DD/MM/AAAA).")
            ws.add_data_validation(dv)
            dv.add(rango)
            for fila in range(2, FILAS + 1):
                ws[f"{letra[nombre]}{fila}"].number_format = "DD/MM/YYYY"
        elif tipo == "número":
            dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="120", allow_blank=True,
                                showErrorMessage=True, errorTitle="Pena no válida", error="Escriba la pena en años (por ejemplo 12.5).")
            ws.add_data_validation(dv)
            dv.add(rango)

    # Instrucciones
    ins = wb.create_sheet("Instrucciones", 0)
    ins.append(["Plantilla de captura — CAUSAS PENALES PPL"])
    ins["A1"].font = Font(bold=True, size=14)
    ins.append(["Llene la hoja 'Captura' (un renglón por persona y causa). No cambie los encabezados."])
    ins.append(["Guárdela como 'CSV UTF-8 (delimitado por comas)' para cargarla en la aplicación."])
    ins.append([])
    ins.append(["Columna", "¿Obligatoria?", "Formato", "Qué capturar"])
    for c in ins[5]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab
    for nombre, oblig, tipo, desc in COLUMNAS:
        ins.append([nombre, oblig, tipo, desc])
    for col, ancho in zip("ABCD", [40, 16, 12, 100]):
        ins.column_dimensions[col].width = ancho
    for fila in ins.iter_rows(min_row=6):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print(f"Plantilla generada: {SALIDA.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
